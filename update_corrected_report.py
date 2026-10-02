"""Update experiment-dependent content while keeping the existing Word layout."""
import csv,json,shutil,datetime
from pathlib import Path
from docx import Document
from PIL import Image
from docx.oxml.ns import qn
from corrected_report_content import discussion_content,pct
ROOT=Path(__file__).resolve().parent
RESULTS=ROOT/'results'/'corrected_group_split_20261002'
REPORT=ROOT/'LAB1_315553010_楊敦傑.docx'
ORDER=('resnet18','resnet50','resnet101')

def replace_text(paragraph,text):
    if paragraph.runs:
        paragraph.runs[0].text=text
        for run in paragraph.runs[1:]:run.text=''
    else:paragraph.add_run(text)

def main():
    if not (RESULTS/'experiment_completed.json').exists():raise RuntimeError('Training and checkpoint audit must finish before updating report')
    summary=json.loads((RESULTS/'comparison_summary.json').read_text(encoding='utf-8'))
    meta=json.loads((RESULTS/'run_metadata.json').read_text(encoding='utf-8'))
    rows=list(csv.DictReader((RESULTS/'combined_epoch_metrics.csv').open(encoding='utf-8')))
    assert len(rows)==36
    latest_backup=ROOT/'report_backups'/('before_corrected_report_update_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+'.docx')
    latest_backup.parent.mkdir(parents=True,exist_ok=True)
    shutil.copy2(REPORT,latest_backup)
    doc=Document(REPORT)
    best=max(ORDER,key=lambda n:summary[n]['best_validation_f1']);high=max(ORDER,key=lambda n:summary[n]['highest_observed_test_accuracy']['accuracy'])
    item=summary[best];test=item['final_test_at_validation_selected_epoch'];highest=summary[high]['highest_observed_test_accuracy']
    for p in doc.paragraphs:
        if p.text.startswith('學號 315553010'):
            replace_text(p,'學號 315553010　姓名 楊敦傑　修正實驗日期 2026 年 10 月 2 日')
        elif p.text.startswith('摘要'):
            replace_text(p,f"摘要　本實驗比較三種 ImageNet 預訓練 ResNet，並修正逐張切分造成的訓練與驗證重複影像洩漏。原始資料全部保留，三種模型以相同分組切分重新訓練。逐輪最高觀察測試準確率為 {high} 第 {highest['epoch']} 輪的 {pct(highest['accuracy'])}；跨模型僅依驗證 F1 選定的權重為 {best} 第 {item['validation_selected_epoch']} 輪，測試準確率 {pct(test['accuracy'])}、肺炎 F1 {test['f1']:.4f}。報告中的曲線、成績及討論均依修正後實驗更新。")
        elif p.text.startswith('使用規格指定的 Kaggle'):
            replace_text(p,'使用指定 Kaggle Chest X Ray Images Pneumonia 第 2 版資料。本機 5856 張影像的相對路徑、大小及 CRC32 全部符合官方 ZIP；資料卡文字的 5863 張與可下載檔案數不同。自訂 ChestXrayDataset 逐張開啟影像、轉 RGB，NORMAL 編碼為 0、PNEUMONIA 為 1。以同一檔名組、檔案 SHA256 及解碼 RGB SHA256 建立連通分組，再依類別與固定種子 315553010，從官方 train 選取約 15% 作驗證；官方 val 保留在驗證集，官方 test 624 張完整保留。肺炎組使用 person 編號加 bacteria／virus 命名範圍，正常組使用 IM／NORMAL2-IM 序號。三模型共用一份經雜湊核對的切分清單，組與相同影像均未跨區；此分組尚未以臨床病患識別碼核實。')
        elif p.text.startswith('預訓練核對'):
            replace_text(p,'預訓練核對：修正實驗的三個交付權重，凍結 conv1 均與對應 torchvision ImageNet 權重逐元素完全相同，分類層均輸出 2 類。紀錄見 results/corrected_group_split_20261002/pretrained_verification.json。')
        elif p.text.startswith('圖 1'):
            replace_text(p,'圖 1　修正實驗結果頁的實際截圖，包含各模型最高觀察測試成績、原始日誌行及驗證選模成績。')
        elif p.text.startswith('三個模型均在 Python'):
            replace_text(p,p.text.split(' 每個模型開始前')[0]+' 每個模型開始前重設相同隨機種子，停用 cuDNN benchmark 並啟用 deterministic 設定；這可減少執行差異，但不保證跨版本或硬體逐位元一致。')
    assert len(doc.tables)==4
    contents=[
        [[s,meta[f'{split}_counts']['0'],meta[f'{split}_counts']['1'],sum(meta[f'{split}_counts'].values())] for s,split in [('訓練','train'),('驗證','validation'),('測試','test')]],
        None,
        [[n,summary[n]['validation_selected_epoch'],f"{summary[n]['best_validation_f1']:.4f}",pct(summary[n]['final_test_at_validation_selected_epoch']['accuracy']),*[f"{summary[n]['final_test_at_validation_selected_epoch'][k]:.4f}" for k in ('precision','recall','f1')]] for n in ORDER],
        [[n,f"{pct(summary[n]['highest_observed_test_accuracy']['accuracy'])}  第 {summary[n]['highest_observed_test_accuracy']['epoch']} 輪",f"{summary[n]['highest_observed_test_f1']['f1']:.4f}  第 {summary[n]['highest_observed_test_f1']['epoch']} 輪",f"{summary[n]['elapsed_seconds']/60:.1f} 分"] for n in ORDER]
    ]
    for table,records in zip(doc.tables,contents):
        if records is None:continue
        for row,values in zip(table.rows[1:],records):
            for cell,value in zip(row.cells,values):replace_text(cell.paragraphs[0],str(value))
    figures=('comparison_results_screenshot.jpg','comparison_accuracy.png','comparison_f1.png','comparison_validation.png','comparison_highest_test_confusion.png','comparison_final_test_confusion.png')
    assert len(doc.inline_shapes)==len(figures)
    for shape,filename in zip(doc.inline_shapes,figures):
        rel=shape._inline.graphic.graphicData.pic.blipFill.blip.get(qn('r:embed'))
        new_rel,_=doc.part.get_or_add_image(str(RESULTS/filename))
        shape._inline.graphic.graphicData.pic.blipFill.blip.set(qn('r:embed'),new_rel)
        if rel!=new_rel and not any(node.get(qn('r:embed'))==rel for node in doc._element.xpath('.//a:blip')):
            doc.part.drop_rel(rel)
        with Image.open(RESULTS/filename) as image:width,height=image.size
        shape.height=round(shape.width*height/width)
        shape._inline.graphic.graphicData.pic.spPr.xfrm.ext.cy=shape.height
    start=next(p._p for p in doc.paragraphs if p.text=='討論' and p.style.name=='Heading 1')
    end=next(p._p for p in doc.paragraphs if p.text=='GitHub 連結' and p.style.name=='Heading 1')
    body=doc._element.body;elements=list(body);a,b=elements.index(start),elements.index(end)
    tail_xml=[e.xml for e in elements[b:]]
    for e in elements[a:b]:body.remove(e)
    for style,text in discussion_content(summary,meta,rows):
        p=doc.add_paragraph(text,style=style);end.addprevious(p._p)
    assert [e.xml for e in list(body)[list(body).index(end):]]==tail_xml
    temp=REPORT.with_suffix('.updated.docx');doc.save(temp)
    verified=Document(temp)
    assert len(verified.tables)==4 and len(verified.inline_shapes)==6
    temp.replace(REPORT)
    print('Updated corrected experiment text/tables/figures/discussion; preserved existing styles and unaffected sections.')
    print(REPORT)
if __name__=='__main__':main()
