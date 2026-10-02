"""Create the editable three-model lab report from recorded D-drive results."""

import csv
import json
from pathlib import Path

from docx import Document
from corrected_report_content import add_discussion
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "results" / "corrected_group_split_20261002"
OUTPUT = ROOT / "LAB1_315553010_楊敦傑.docx"
ORDER = ("resnet18", "resnet50", "resnet101")


def load_data():
    summary = json.loads((RESULTS / "comparison_summary.json").read_text(encoding="utf-8"))
    metadata = json.loads((RESULTS / "run_metadata.json").read_text(encoding="utf-8"))
    with (RESULTS / "combined_epoch_metrics.csv").open(newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    if set(summary) != set(ORDER) or len(rows) != 36:
        raise ValueError("Report requires all three complete 12-epoch runs")
    return summary, metadata, rows


def set_font(style, size, bold=False):
    style.font.name = "Microsoft JhengHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")


def cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def cell_border(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        borders.append(el)


def cell_margin(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{m}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)


def add_table(doc, headings, records, widths):
    table = doc.add_table(rows=1, cols=len(headings))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, heading in enumerate(headings):
        table.columns[i].width = Inches(widths[i])
        table.rows[0].cells[i].text = heading
    for record in records:
        row = table.add_row()
        for i, value in enumerate(record):
            row.cells[i].text = str(value)
    for ri, row in enumerate(table.rows):
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_border(cell)
            cell_margin(cell)
            if ri == 0:
                cell_shading(cell, "DCEAF1")
            elif ri % 2 == 0:
                cell_shading(cell, "F7F9FA")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.name = "Microsoft JhengHei"
                    run.font.size = Pt(8.5)
                    run.font.bold = ri == 0
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft JhengHei")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_body(doc, text):
    p = doc.add_paragraph(text, style="Normal")
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(text, style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_figure(doc, path, caption, width=6.4):
    if not path.exists():
        raise FileNotFoundError(path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption)


def pct(value):
    return f"{100 * value:.2f}%"


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def build():
    if not (RESULTS / "experiment_completed.json").exists():
        raise RuntimeError("Complete corrected training and audit before building report")
    summary, metadata, rows = load_data()
    chosen_name = max(ORDER, key=lambda name: summary[name]["best_validation_f1"])
    highest_name = max(ORDER, key=lambda name: summary[name]["highest_observed_test_accuracy"]["accuracy"])
    chosen = summary[chosen_name]
    highest = summary[highest_name]["highest_observed_test_accuracy"]
    n_train = sum(metadata["train_counts"].values())
    n_val = sum(metadata["validation_counts"].values())
    n_test = sum(metadata["test_counts"].values())
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = Mm(22)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(21)
    section.right_margin = Mm(21)
    section.header_distance = Mm(10)
    section.footer_distance = Mm(11)

    styles = doc.styles
    set_font(styles["Normal"], 10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.28
    set_font(styles["Title"], 20, True)
    styles["Title"].paragraph_format.space_after = Pt(8)
    set_font(styles["Subtitle"], 11)
    styles["Subtitle"].paragraph_format.space_after = Pt(10)
    set_font(styles["Heading 1"], 14, True)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)
    styles["Heading 1"].paragraph_format.keep_with_next = True
    set_font(styles["Heading 2"], 11.5, True)
    styles["Heading 2"].paragraph_format.space_before = Pt(10)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)
    styles["Heading 2"].paragraph_format.keep_with_next = True
    set_font(styles["Caption"], 9)
    styles["Caption"].paragraph_format.space_after = Pt(9)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.add_run("315553010  楊敦傑")
    for run in header.runs:
        run.font.name = "Microsoft JhengHei"
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0, 0, 0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_page_number(footer)

    doc.add_paragraph("胸部 X 光影像肺炎分類實驗報告", style="Title")
    doc.add_paragraph("ResNet18 ResNet50 與 ResNet101 比較", style="Subtitle")
    add_body(doc, "學號 315553010　姓名 楊敦傑　修正實驗日期 2026 年 10 月 2 日")
    selected_test = chosen["final_test_at_validation_selected_epoch"]
    add_body(doc, f"摘要　本實驗以自訂資料載入器比較三種 ImageNet 預訓練 ResNet。"
             f"逐輪最高觀察測試準確率為 {highest_name} 第 {highest['epoch']} 輪的 {pct(highest['accuracy'])}。"
             f"僅依驗證 F1 選定的交付權重為 {chosen_name} 第 {chosen['validation_selected_epoch']} 輪，"
             f"其測試準確率為 {pct(selected_test['accuracy'])}、肺炎 F1 為 {selected_test['f1']:.4f}。"
             "最高觀察值與驗證選模後的結果分別呈現，避免混淆測試集用途。")

    doc.add_heading("研究目標", level=1)
    add_body(doc, "胸部 X 光影像分類的目標為區分正常與肺炎。本作業要求自訂影像資料載入、至少兩種 ResNet、逐輪準確率及 F1 曲線，並以混淆矩陣呈現測試結果。ResNet 的殘差捷徑使較深網路更容易最佳化。本實驗以 ResNet18、ResNet50 與 ResNet101 比較網路深度、計算成本及測試表現。")

    doc.add_heading("實驗設定", level=1)
    doc.add_heading("資料載入與切分", level=2)
    add_body(doc, "使用規格指定的 Kaggle Chest X Ray Images Pneumonia 資料集。自訂 ChestXrayDataset 逐筆開啟 JPEG 或 PNG，轉為 RGB，將 NORMAL 編碼為 0、PNEUMONIA 編碼為 1。本機 5856 張影像與官方第 2 版 ZIP 路徑、大小及 CRC32 全部相符。以檔名組、檔案 SHA256 與解碼 RGB SHA256 建立連通分組，再依類別及固定種子 315553010，從官方 train 選取約 15% 作驗證；官方 val 保留在驗證集，官方 test 624 張完整保留。肺炎組使用 person 編號加 bacteria／virus 命名範圍，正常組使用 IM／NORMAL2-IM 序號。三次訓練共用切分清單，分組與相同影像均未跨區，病患臨床識別碼尚未核實。")
    add_table(doc, ["資料切分", "NORMAL", "PNEUMONIA", "總數"], [
        ["訓練", metadata["train_counts"]["0"], metadata["train_counts"]["1"], n_train],
        ["驗證", metadata["validation_counts"]["0"], metadata["validation_counts"]["1"], n_val],
        ["測試", metadata["test_counts"]["0"], metadata["test_counts"]["1"], n_test],
    ], [1.4, 1.45, 1.55, 1.2])
    add_body(doc, "訓練影像縮放至 224 × 224，再施加 ±7° 旋轉、最多 3% 平移、0.95 至 1.05 倍縮放，以及最多 10% 的亮度與對比變動。驗證與測試影像只縮放，不做隨機增強。所有影像以 ImageNet 均值及標準差正規化。")
    doc.add_heading("模型架構與預訓練", level=2)
    add_body(doc, "ResNet18 使用 basic block，ResNet50 與 ResNet101 使用 bottleneck block。三者均由 torchvision 載入 ImageNet 預訓練權重，再重新初始化二類分類層。訓練時凍結前段卷積、layer1 與 layer2，微調 layer3、layer4 與分類層。ResNet101 讓深度相對 ResNet18 大幅增加，可檢驗較深模型是否帶來實際收益；較深不保證較準。")
    add_table(doc, ["模型", "預訓練權重", "分類層", "總參數", "可訓練參數"], [
        ["ResNet18", "IMAGENET1K_V1", "512 → 2", "11,177,538", "10,494,466"],
        ["ResNet50", "IMAGENET1K_V2", "2048 → 2", "23,512,130", "22,067,202"],
        ["ResNet101", "IMAGENET1K_V2", "2048 → 2", "42,504,258", "41,059,330"],
    ], [1.1, 1.65, 0.9, 1.2, 1.35])
    add_body(doc, "預訓練核對：三個交付權重的凍結 conv1 參數，均與對應 torchvision ImageNet 權重逐元素完全相同；分類層均輸出 2 類。檢查紀錄見 results/corrected_group_split_20261002/pretrained_verification.json。")
    doc.add_heading("訓練方法與環境", level=2)
    add_body(doc, f"三個模型均在 Python 內建 venv 與 {metadata['gpu']} 上訓練，"
             f"使用 PyTorch {metadata['torch']}、torchvision {metadata['torchvision']}、CUDA {metadata['cuda']}。"
             "每個模型訓練 12 輪、batch size 32，以訓練類別比例的倒數計算加權交叉熵，"
             "使用 AdamW 學習率 1 × 10⁻⁴、weight decay 1 × 10⁻⁴、cosine 學習率衰減與 CUDA 混合精度。"
             "每輪記錄訓練、驗證及測試指標，僅以驗證集肺炎 F1 選擇保存權重。每模型開始前重設種子，停用 cuDNN benchmark 並啟用 deterministic 設定，仍不保證跨版本或硬體逐位元相同。")

    doc.add_heading("實驗結果", level=1)
    doc.add_heading("驗證選模後的測試結果", level=2)
    selected_records = []
    for name in ORDER:
        item = summary[name]
        score = item["final_test_at_validation_selected_epoch"]
        selected_records.append([name, item["validation_selected_epoch"],
                                 f"{item['best_validation_f1']:.4f}", pct(score["accuracy"]),
                                 f"{score['precision']:.4f}", f"{score['recall']:.4f}", f"{score['f1']:.4f}"])
    add_table(doc, ["模型", "輪次", "Val F1", "Test Acc", "Precision", "Recall", "Test F1"],
              selected_records, [1.0, 0.55, 0.8, 1.0, 1.05, 0.85, 0.9])
    doc.add_heading("逐輪最高觀察測試成績", level=2)
    observed_records = []
    for name in ORDER:
        item = summary[name]
        observed_records.append([name,
                                 f"{pct(item['highest_observed_test_accuracy']['accuracy'])}  第 {item['highest_observed_test_accuracy']['epoch']} 輪",
                                 f"{item['highest_observed_test_f1']['f1']:.4f}  第 {item['highest_observed_test_f1']['epoch']} 輪",
                                 f"{item['elapsed_seconds']/60:.1f} 分"])
    add_table(doc, ["模型", "最高 Test Acc", "最高 Test F1", "訓練耗時"],
              observed_records, [1.1, 2.05, 2.0, 1.15])
    add_figure(doc, RESULTS / "comparison_results_screenshot.jpg",
               "圖 1　修正實驗結果頁的實際截圖，包含最高觀察測試成績、原始日誌行及驗證選模成績。")
    add_body(doc, "上表最高準確率與最高 F1 可能出自不同輪次。逐輪測試曲線依作業要求保留，但未用於權重或超參數選擇。")
    add_figure(doc, RESULTS / "comparison_accuracy.png",
               "圖 2　三種 ResNet 每輪訓練與測試準確率。實線為測試，虛線為訓練。")
    add_figure(doc, RESULTS / "comparison_f1.png",
               "圖 3　三種 ResNet 每輪訓練與測試肺炎 F1。實線為測試，虛線為訓練。")
    add_figure(doc, RESULTS / "comparison_validation.png",
               "圖 4　驗證集準確率與 F1 曲線，補充呈現權重選擇依據。")
    doc.add_heading("混淆矩陣", level=2)
    add_figure(doc, RESULTS / "comparison_highest_test_confusion.png",
               "圖 5　各模型最高觀察測試準確率輪次的混淆矩陣。橫軸為預測，縱軸為實際。")
    add_figure(doc, RESULTS / "comparison_final_test_confusion.png",
               "圖 6　各模型依驗證 F1 選定權重後的正式測試混淆矩陣。")

    add_discussion(doc, summary, metadata, rows)

    doc.add_heading("GitHub 連結", level=1)
    add_body(doc, "待本人確認內容、建立 Git 儲存庫並自行推送後填入：____________________________")
    doc.add_heading("參考資料", level=1)
    add_body(doc, "1. He 等人，Deep Residual Learning for Image Recognition，2015。https://arxiv.org/abs/1512.03385")
    add_body(doc, "2. Torchvision ResNet API。https://docs.pytorch.org/vision/stable/models/resnet.html")
    add_body(doc, "3. Kaggle Chest X Ray Images Pneumonia。https://www.kaggle.com/datasets/paultimothymooney/chest-xray-pneumonia")

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
