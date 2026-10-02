import datetime,hashlib,json,zipfile
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
OUT=ROOT/'results'/'corrected_group_split_20261002'
summary=json.loads((OUT/'summary.json').read_text(encoding='utf-8'))
report=ROOT/'LAB1_315553010_楊敦傑.docx'
doc=Document(report)
with zipfile.ZipFile(report) as z:assert z.testzip() is None
assert len(doc.tables)==4 and len(doc.inline_shapes)==6
for row,name in zip(doc.tables[2].rows[1:],('resnet18','resnet50','resnet101')):
    item=summary[name];test=item['final_test_at_validation_selected_epoch']
    expected=[name,str(item['validation_selected_epoch']),f"{item['best_validation_f1']:.4f}",f"{test['accuracy']*100:.2f}%",*[f"{test[k]:.4f}" for k in ('precision','recall','f1')]]
    assert [c.text for c in row.cells]==expected
figures=('comparison_results_screenshot.jpg','comparison_accuracy.png','comparison_f1.png','comparison_validation.png','comparison_highest_test_confusion.png','comparison_final_test_confusion.png')
image_types=[]
for shape,name in zip(doc.inline_shapes,figures):
    rel=shape._inline.graphic.graphicData.pic.blipFill.blip.get(qn('r:embed'))
    part=doc.part.related_parts[rel]
    assert part.blob==(OUT/name).read_bytes()
    assert part.content_type==('image/jpeg' if name.endswith('.jpg') else 'image/png')
    image_types.append(part.content_type)
text='\n'.join(p.text for p in doc.paragraphs)
assert '30 組' in text and '第 2 版' in text and '總錯誤從 77 變為 67' in text
assert '本次三者耗時差距很小' in text and '像素雜湊重疊均為 0' in text
assert 'Github' not in text or 'GitHub 連結' in text
check={'docx_zip_crc_ok':True,'tables':4,'images':6,'all_selected_result_table_values_match':True,
       'all_embedded_figures_match_corrected_sources':True,'image_content_types':image_types,
       'discussion_matches_corrected_data':True,'git_not_initialized':not (ROOT/'.git').exists(),
       'automatic_pagination_checked':False}
(ROOT/'report_qa').mkdir(exist_ok=True)
(ROOT/'report_qa'/'corrected_run_content_check.json').write_text(json.dumps(check,indent=2),encoding='utf-8')
proof=dict(check,report_sha256=hashlib.sha256(report.read_bytes()).hexdigest(),review_date=datetime.date.today().isoformat())
(OUT/'report_verification.json').write_text(json.dumps(proof,indent=2),encoding='utf-8')
print(json.dumps(check,indent=2))

