"""Update only the discussion using the completed corrected experiment."""
import csv,datetime,json,shutil
from pathlib import Path
from docx import Document
from corrected_report_content import discussion_content
ROOT=Path(__file__).resolve().parent
RESULTS=ROOT/'results'/'corrected_group_split_20261002'
REPORT=ROOT/'LAB1_315553010_楊敦傑.docx'
def main():
    if not (RESULTS/'audit_corrected.json').exists():raise RuntimeError('Corrected checkpoint audit is required')
    summary=json.loads((RESULTS/'comparison_summary.json').read_text(encoding='utf-8'))
    meta=json.loads((RESULTS/'run_metadata.json').read_text(encoding='utf-8'))
    rows=list(csv.DictReader((RESULTS/'combined_epoch_metrics.csv').open(encoding='utf-8')))
    doc=Document(REPORT);body=doc._element.body
    start=next(p._p for p in doc.paragraphs if p.text=='討論' and p.style.name=='Heading 1')
    end=next(p._p for p in doc.paragraphs if p.text=='GitHub 連結' and p.style.name=='Heading 1')
    elements=list(body);a,b=elements.index(start),elements.index(end)
    head=[e.xml for e in elements[:a]];tail=[e.xml for e in elements[b:]]
    backup=ROOT/'report_backups'/('before_corrected_discussion_'+datetime.datetime.now().strftime('%Y%m%d_%H%M%S')+'.docx')
    shutil.copy2(REPORT,backup)
    for element in elements[a:b]:body.remove(element)
    inserted=[]
    for style,text in discussion_content(summary,meta,rows):
        p=doc.add_paragraph(text,style=style);end.addprevious(p._p);inserted.append(p._p)
    now=list(body)
    assert [e.xml for e in now[:a]]==head and [e.xml for e in now[a+len(inserted):]]==tail
    doc.save(REPORT)
    print('Discussion updated; all other body XML unchanged. Backup:',backup)
if __name__=='__main__':main()
